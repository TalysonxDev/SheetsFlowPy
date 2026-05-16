
import pandas as pd
import openpyxl

# ──────────────────────────────────────────────────────────────────────────────
# Caminhos dos arquivos
# ──────────────────────────────────────────────────────────────────────────────

ARQUIVO_FLUXO       = r"C:\Users\Talyson Negrao\OneDrive\Documentos\Sitran\Projeto Automação Fluxo\Fluxo.xlsx"
ARQUIVO_FLUXO_DIARIO = r"C:\Users\Talyson Negrao\OneDrive\Documentos\Sitran\Projeto Automação Fluxo\FluxoDiario.xlsx"

# ──────────────────────────────────────────────────────────────────────────────
# Índices de colunas no FluxoDiario (base 0)
# ──────────────────────────────────────────────────────────────────────────────

COL_GRIN      = 3   # Coluna D — código GRIN da faixa
COL_ENDERECO  = 4   # Coluna E — endereço da faixa


# ══════════════════════════════════════════════════════════════════════════════
# PASSO 1 — Leitura do FluxoDiario.xlsx
# ══════════════════════════════════════════════════════════════════════════════

def ler_fluxo_diario(caminho: str) -> tuple[dict, int]:
    
    xl   = pd.ExcelFile(caminho)
    abas = [aba for aba in xl.sheet_names if aba != "Page 16"]

    registros = []  # Lista de tuplas (grin_int, dia, volume)

    for aba in abas:
        # Lê a aba sem cabeçalho definido, pois o layout varia entre abas.
        df = pd.read_excel(caminho, sheet_name=aba, header=None)

        # ── Detecção do cabeçalho de dias ────────────────────────────────────
        # A linha de cabeçalho é identificada por conter 5 ou mais valores parecidos com dias
        idx_header = None
        for i, row in df.iterrows():
            dias_encontrados = []
            for val in row:
                try:
                    d = int(float(val))
                    if 1 <= d <= 31:
                        dias_encontrados.append(d)
                except (ValueError, TypeError):
                    pass  # Ignora células não numéricas

            if len(dias_encontrados) >= 5:
                idx_header = i
                break  # Linha de cabeçalho localizada

        if idx_header is None:
            print(f"  ⚠ Aba '{aba}': cabeçalho não encontrado, pulando.")
            continue

        # Linha de cabeçalho e dados abaixo dela
        header_row = df.iloc[idx_header]
        dados      = df.iloc[idx_header + 1:].reset_index(drop=True)

        # ── Mapeamento dia → índice de coluna ────────────────────────────────
        dia_para_col = {}
        for col_idx, val in header_row.items():
            try:
                d = int(float(val))
                if 1 <= d <= 31:
                    dia_para_col[d] = col_idx
            except (ValueError, TypeError):
                pass

        # ── Coleta de registros válidos ───────────────────────────────────────
        count = 0
        for _, row in dados.iterrows():
            grin     = row.iloc[COL_GRIN]     if COL_GRIN     < len(row) else None
            endereco = str(row.iloc[COL_ENDERECO]).strip() if COL_ENDERECO < len(row) else ''

            # Ignora linhas sem GRIN válido
            if pd.isna(grin):
                continue
            try:
                grin_int = int(float(grin))
            except (ValueError, TypeError):
                continue
            if grin_int == 0:
                continue

            # Ignora faixas sem endereço
            if endereco in ('0', '', 'nan'):
                continue

            # Coleta o volume de cada dia mapeado
            for dia, col_idx in dia_para_col.items():
                volume = row.iloc[col_idx] if col_idx < len(row) else None
                if pd.notna(volume):
                    try:
                        registros.append((grin_int, dia, float(volume)))
                        count += 1
                    except (ValueError, TypeError):
                        pass

        print(f"  ✓ Aba '{aba}': header L{idx_header + 1}, {len(dia_para_col)} dias, {count} registros")

    # ── Validações finais ─────────────────────────────────────────────────────
    if not registros:
        raise ValueError("Nenhum dado encontrado no FluxoDiario.xlsx")

    dias_com_dados = {dia for _, dia, vol in registros if vol > 0}
    if not dias_com_dados:
        raise ValueError("Nenhum dia com valores > 0 encontrado")

    ultimo_dia = max(dias_com_dados)
    print(f"\nÚltimo dia com dados: {ultimo_dia}")

    # Monta o mapa {GRIN: volume} apenas para o último dia
    mapa = {grin: vol for grin, dia, vol in registros if dia == ultimo_dia}
    print(f"Total de GRINs com dados no dia {ultimo_dia}: {len(mapa)}")

    return mapa, ultimo_dia


# ══════════════════════════════════════════════════════════════════════════════
# PASSO 2 — Preenchimento do Fluxo.xlsx
# Localiza a coluna do último dia na linha 3 e aplica as regras de status
# a partir da linha 4 em diante.
# ══════════════════════════════════════════════════════════════════════════════

mapa, ultimo_dia = ler_fluxo_diario(ARQUIVO_FLUXO_DIARIO)

wb = openpyxl.load_workbook(ARQUIVO_FLUXO)
ws = wb.active

# ── Localização das colunas de dia no Fluxo.xlsx ─────────────────────────────
# Os dias ficam na linha 3. A busca é limitada à coluna 50 para evitar
# capturar colunas auxiliares que possam conter números coincidentes.

col_dia_num      = None   # Número da coluna do último dia
col_dia_letra    = None   # Letra da coluna do último dia (para log)
col_anterior_num = None   # Número da coluna do dia anterior

for cell in ws[3]:
    if cell.column > 50:
        break
    try:
        val = int(cell.value)
        if val == ultimo_dia:
            col_dia_num   = cell.column
            col_dia_letra = cell.column_letter
        if val == ultimo_dia - 1:
            col_anterior_num = cell.column
    except (TypeError, ValueError):
        pass

if col_dia_num is None:
    raise ValueError(f"Coluna do dia {ultimo_dia} não encontrada em Fluxo.xlsx (linha 3)")

print(f"Coluna do dia {ultimo_dia}: {col_dia_letra} ({col_dia_num})")
print(f"Coluna do dia anterior ({ultimo_dia - 1}): {col_anterior_num}")
print("Preenchendo...")

# ── Aplicação das regras de status ───────────────────────────────────────────
# A coluna C (índice 3) do Fluxo.xlsx contém o código GRIN de cada faixa.
# Regras de negócio:
#   Volume >= 13 → OK   se anterior for OK/F/V/O/AF
#               → LR   se anterior for LR
#               → repete anterior nos demais casos
#   Volume <  13 → F    se anterior for OK/F
#               → repete anterior nos demais casos

COL_GRIN_FLUXO = 3  # Coluna C no Fluxo.xlsx

atualizados = 0
ignorados   = 0

for row in ws.iter_rows(min_row=4):
    grin_cell = row[COL_GRIN_FLUXO - 1]
    try:
        grin = int(float(grin_cell.value))
    except (TypeError, ValueError):
        continue

    volume = mapa.get(grin)

    if volume is not None:
        celula_dia = row[col_dia_num - 1]
        ontem_raw  = row[col_anterior_num - 1].value if col_anterior_num else None
        ontem      = str(ontem_raw).strip().upper()  if ontem_raw is not None else ""

        if volume >= 13:
            if ontem in ("OK", "F", "V", "O", "AF"):
                celula_dia.value = "OK"
            elif ontem == "LR":
                celula_dia.value = "LR"
            else:
                celula_dia.value = ontem_raw
        else:
            if ontem in ("OK", "F"):
                celula_dia.value = "F"
            else:
                celula_dia.value = ontem_raw

        atualizados += 1
    else:
        ignorados += 1

wb.save(ARQUIVO_FLUXO)

print(f"\n✅ Concluído!")
print(f"   Dia {ultimo_dia} preenchido com sucesso!")
print(f"   Faixas atualizadas : {atualizados}")
print(f"   Faixas ignoradas   : {ignorados}  (sem endereço ou ausentes no FluxoDiario)")